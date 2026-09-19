"""Assemble the complete dissertation into the Pan-Atlantic MSc template.

The template is used as the base document rather than imitated, so the styles, the table
style and the section setup are the university's own. Its placeholder body is removed and
the real content written back through the same styles: Heading 1 for chapter titles,
Heading 2 and 3 for sections, Section Title for front matter, Table/Figure for table notes
and figure captions, and APA Report for tables.

Front matter is numbered in lower-case roman and the body in arabic restarting at one,
which needs raw XML because python-docx exposes neither page-number format nor fields.
"""
from __future__ import annotations
import re
import sys
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt

D = Path("dissertation")
TEMPLATE = D / "PAU_TEMPLATE.docx"
OUT = D / "Maduechesi_25120133019_MSc_Data_Science_Project.docx"

STUDENT = "Maduechesi Chidiebere Jennifer"
MATRIC = "25120133019"
DEPT = "Computer and Information Sciences Department"
SCHOOL = "School of Science and Technology"
UNI = "Pan-Atlantic University"
DATE = "September 2026"
SUPERVISORS = [("Mr Solomon Alile", "Supervisor 1"), ("Dr Adubi", "Supervisor 2")]

CHAPTERS = [
    ("Chapter One: Introduction", "chapter1_introduction_v2.md"),
    ("Chapter Two: Literature Review", "chapter2_literature_review.md"),
    ("Chapter Three: Methodology", "chapter3_methodology.md"),
    ("Chapter Four: Results Discussion", "chapter4_results.md"),
    ("Chapter Five: Summary, Conclusions and Recommendations", "chapter5_conclusions.md"),
]
FIGURE_WIDTH = Inches(5.9)


# --------------------------------------------------------------------------- xml helpers
def field(paragraph, instr: str) -> None:
    """Insert a Word field, which python-docx has no API for."""
    run = paragraph.add_run()
    for kind, text in (("begin", None), ("instrText", instr), ("end", None)):
        el = OxmlElement(f"w:fld{kind.capitalize()}" if kind != "instrText" else "w:instrText")
        if kind == "instrText":
            el.set(qn("xml:space"), "preserve")
            el.text = text
        else:
            el = OxmlElement("w:fldChar")
            el.set(qn("w:fldCharType"), kind)
        run._r.append(el)


def page_numbering(section, fmt: str, start: int | None, title_page: bool) -> None:
    """Set the page-number format on the section's existing pgNumType.

    Appending a second one leaves two in the same sectPr, and appending it after
    <w:titlePg/> also puts it out of schema order, so Word ignores it and the front
    matter comes out in arabic like the body.
    """
    sectPr = section._sectPr
    pgnum = sectPr.find(qn("w:pgNumType"))
    if pgnum is None:
        pgnum = OxmlElement("w:pgNumType")
        anchor = sectPr.find(qn("w:titlePg")) or sectPr.find(qn("w:cols"))
        (sectPr.insert(list(sectPr).index(anchor), pgnum) if anchor is not None
         else sectPr.append(pgnum))
    pgnum.set(qn("w:fmt"), fmt)
    if start is not None:
        pgnum.set(qn("w:start"), str(start))
    # titlePg suppresses the number on the first page of the section, which is wanted on
    # the title page and not on the first page of the body.
    existing = sectPr.find(qn("w:titlePg"))
    if title_page and existing is None:
        sectPr.append(OxmlElement("w:titlePg"))
    elif not title_page and existing is not None:
        sectPr.remove(existing)


def footer_page_number(section) -> None:
    section.footer.is_linked_to_previous = False
    p = section.footer.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    field(p, "PAGE")


def runs_with_markup(paragraph, text: str) -> None:
    """Render **bold** and *italic* spans, leaving everything else plain."""
    for part in re.split(r"(\*\*[^*]+\*\*|\*[^*]+\*)", text):
        if not part:
            continue
        if part.startswith("**") and part.endswith("**"):
            paragraph.add_run(part[2:-2]).bold = True
        elif part.startswith("*") and part.endswith("*"):
            paragraph.add_run(part[1:-1]).italic = True
        else:
            paragraph.add_run(part)


BODY_STYLES = {None, "Normal"}


def para(doc, text: str = "", style: str | None = None, align=None, indent_first=None):
    p = doc.add_paragraph(style=style)
    if align is not None:
        p.alignment = align
    elif style in BODY_STYLES:
        # Body prose is justified. Headings, captions and the centred blocks set their
        # own alignment and are left alone.
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    if indent_first is not None:
        p.paragraph_format.first_line_indent = indent_first
    if text:
        runs_with_markup(p, text)
    return p


# ------------------------------------------------------------------------------- tables
def add_table(doc, rows: list[list[str]]) -> None:
    t = doc.add_table(rows=len(rows), cols=len(rows[0]))
    try:
        t.style = doc.styles["APA Report"]
    except KeyError:
        t.style = doc.styles["Table Grid"]
    for r, cells in enumerate(rows):
        for c, value in enumerate(cells):
            cell = t.cell(r, c)
            cell.text = ""
            p = cell.paragraphs[0]
            p.paragraph_format.first_line_indent = Pt(0)
            p.paragraph_format.space_before = Pt(2)
            p.paragraph_format.space_after = Pt(2)
            pf = p.paragraph_format._element.get_or_add_pPr()
            spacing = OxmlElement("w:spacing")
            spacing.set(qn("w:line"), "240")
            spacing.set(qn("w:lineRule"), "auto")
            pf.append(spacing)
            runs_with_markup(p, value)
            for run in p.runs:
                run.font.size = Pt(10)
                if r == 0:
                    run.bold = True


# ------------------------------------------------------------------------- front matter
def front_matter(doc, fm: dict) -> None:
    def centred(text, style="Title 2", blank_after=0):
        para(doc, text, style=style, align=WD_ALIGN_PARAGRAPH.CENTER)
        for _ in range(blank_after):
            para(doc)

    for second_page in (False, True):
        para(doc, fm["title"], style="Title", align=WD_ALIGN_PARAGRAPH.CENTER)
        para(doc); para(doc)
        centred("By", blank_after=1)
        centred(STUDENT)
        centred(MATRIC, blank_after=2)
        if second_page:
            centred(f"A project submitted to the {SCHOOL}, {UNI}")
            centred("in partial fulfillment of the requirements for the award of the degree of")
            centred("Master of Science (Data Science)", blank_after=2)
        else:
            centred(DEPT)
            centred(SCHOOL)
            centred(UNI, blank_after=2)
        centred(DATE)
        doc.add_page_break()

    para(doc, "ABSTRACT", style="Section Title")
    para(doc, fm["abstract"], indent_first=Inches(0.5))
    para(doc)
    para(doc, fm["keywords"], indent_first=Pt(0))
    doc.add_page_break()

    for heading, key in (("ACKNOWLEDGEMENTS", "acknowledgements"), ("DEDICATION", "dedication")):
        para(doc, heading, style="Section Title")
        for block in fm[key].split("\n\n"):
            para(doc, block.strip(), indent_first=Inches(0.5))
        doc.add_page_break()

    para(doc, "STUDENT'S DECLARATION", style="Section Title")
    para(doc, fm["declaration"], indent_first=Inches(0.5))
    for _ in range(3):
        para(doc)
    for line in ("__________________________", STUDENT, MATRIC):
        para(doc, line, align=WD_ALIGN_PARAGRAPH.CENTER, indent_first=Pt(0))
    para(doc)
    para(doc, "Date: ______________", align=WD_ALIGN_PARAGRAPH.CENTER, indent_first=Pt(0))
    doc.add_page_break()

    para(doc, "CERTIFICATION", style="Section Title")
    para(doc, fm["certification"], indent_first=Inches(0.5))
    for name, role in SUPERVISORS:
        for _ in range(2):
            para(doc)
        for line in ("________________________________", f"{role}: {name}", SCHOOL, UNI,
                     "Lagos, Nigeria"):
            para(doc, line, align=WD_ALIGN_PARAGRAPH.CENTER, indent_first=Pt(0))
    para(doc)
    para(doc, "Date: ______________", align=WD_ALIGN_PARAGRAPH.CENTER, indent_first=Pt(0))
    doc.add_page_break()

    listings = fm.get("listings") or {}
    for heading, key in (("TABLE OF CONTENTS", "contents"),
                         ("LIST OF TABLES", "tables"),
                         ("LIST OF FIGURES", "figures")):
        para(doc, heading, style="Section Title")
        rows = listings.get(key)
        if rows:
            add_listing(doc, rows)
        else:
            para(doc, "[This listing is generated when the document is assembled.]",
                 indent_first=Pt(0))
        doc.add_page_break()


def add_listing(doc, rows: list[tuple[int, str, str]]) -> None:
    """Write one listing line per entry: text, dot leader, page number.

    Word sets a right tab with a dot leader at the text margin, which is what puts the
    page number at the right edge with dots running up to it.
    """
    from docx.enum.text import WD_TAB_ALIGNMENT, WD_TAB_LEADER
    right_edge = Inches(6.5)
    for level, text, page in rows:
        p = doc.add_paragraph()
        pf = p.paragraph_format
        pf.first_line_indent = Pt(0)
        pf.left_indent = Inches(0.25 * level)
        pf.space_after = Pt(0)
        _set_single(p)
        pf.tab_stops.add_tab_stop(right_edge - Inches(0.25 * level),
                                  WD_TAB_ALIGNMENT.RIGHT, WD_TAB_LEADER.DOTS)
        run = p.add_run(text)
        if level == 0:
            run.bold = True
        p.add_run("\t" + page)


def _set_single(paragraph) -> None:
    pf = paragraph.paragraph_format._element.get_or_add_pPr()
    spacing = OxmlElement("w:spacing")
    spacing.set(qn("w:line"), "276")
    spacing.set(qn("w:lineRule"), "auto")
    spacing.set(qn("w:after"), "0")
    pf.append(spacing)


# ------------------------------------------------------------------------------ chapters
TABLE_NUM = re.compile(r"^Table (\d+)$")
# Algorithms carry their own caption series, so they never disturb table numbering.
ALGO_NUM = re.compile(r"^Algorithm (\d+)$")
FIGURE_IMG = re.compile(r"^!\[[^\]]*\]\(([^)]+)\)$")


def add_markdown(doc, path: Path, heading_text: str) -> None:
    lines = path.read_text(encoding="utf-8").split("\n")
    para(doc, heading_text, style="Heading 1")
    buffer: list[list[str]] = []
    i = 0
    while i < len(lines):
        s = lines[i].strip()
        if s.startswith("|"):
            cells = [c.strip() for c in s.strip("|").split("|")]
            if not all(set(c) <= set("-: ") for c in cells):
                buffer.append(cells)
            i += 1
            continue
        if buffer:
            add_table(doc, buffer)
            buffer = []
        if not s:
            i += 1
            continue
        if s.startswith("# "):
            i += 1
            continue
        if s.startswith("### "):
            para(doc, s[4:], style="Heading 3")
        elif s.startswith("## "):
            para(doc, s[3:], style="Heading 2")
        elif TABLE_NUM.match(s) or ALGO_NUM.match(s):
            para(doc, s, style="No Spacing", indent_first=Pt(0))
        elif FIGURE_IMG.match(s):
            target = (path.parent / FIGURE_IMG.match(s).group(1)).resolve()
            p = para(doc, style="Table/Figure", align=WD_ALIGN_PARAGRAPH.CENTER,
                     indent_first=Pt(0))
            p.add_run().add_picture(str(target), width=FIGURE_WIDTH)
        elif s.startswith("*Figure "):
            caption = s.replace("*", "")
            cp = para(doc, style="Table/Figure", align=WD_ALIGN_PARAGRAPH.CENTER,
                      indent_first=Pt(0))
            run = cp.add_run(caption)
            run.italic = False
            run.bold = False
            run.font.size = Pt(10)
        elif s.startswith("*Note*"):
            para(doc, s, style="Table/Figure", indent_first=Pt(0))
        elif s.startswith("*") and s.endswith("*") and s.count("*") == 2:
            para(doc, s, style="Table/Figure", indent_first=Pt(0))
        else:
            para(doc, s, indent_first=Inches(0.5))
        i += 1
    if buffer:
        add_table(doc, buffer)


def references(doc) -> None:
    para(doc, "References", style="Heading 1")
    text = (D / "references_consolidated.md").read_text(encoding="utf-8")
    start = re.compile(r"^(?:[^\Wa-z\d_]|(?:d[aeiou]|van|von|del|della|dos|la|le|ten|ter)\s)"
                       r"(?P<a>[^(]{2,220}?)\s*\((?:\d{4}[a-z]?|n\.d\.)\)")
    for line in text.split("\n"):
        s = " ".join(line.split())
        if not s or s.startswith("#"):
            continue
        m = start.match(s)
        entry = bool(m) and (re.search(r",\s*[A-Z]\.", m.group("a"))
                             or m.group("a").rstrip().endswith("."))
        p = para(doc, s, indent_first=Pt(0))
        if entry:
            p.paragraph_format.left_indent = Inches(0.5)
            p.paragraph_format.first_line_indent = Inches(-0.5)


def appendices(doc) -> None:
    para(doc, "Appendices", style="Heading 1")
    para(doc, "Appendix A: Data and Modelling Pipeline",
         style="Heading 2 for Appendix" if "Heading 2 for Appendix"
         in [s.name for s in doc.styles] else "Heading 2")
    para(doc, "The flowchart below traces the study end to end, from the two supplied "
              "incident files to the deployed interface, with the record counts surviving "
              "each stage. The verification box marks the checking scripts described in "
              "Chapter Three, which re-derive every numeric claim in Chapters Three to "
              "Five from the stored artefacts.", indent_first=Inches(0.5))
    p = para(doc, style="Table/Figure", align=WD_ALIGN_PARAGRAPH.CENTER, indent_first=Pt(0))
    p.add_run().add_picture(str((D / "figure8_pipeline.png").resolve()), width=Inches(5.4))
    para(doc, "*Figure 8*. The data and modelling pipeline.", style="Table/Figure",
         indent_first=Pt(0))

    appendix_b(doc)


def appendix_b(doc) -> None:
    """Render the review matrix. Shared by the full build and the Chapter Two review file."""
    para(doc, "Appendix B: Matrix of Reviewed Literature",
         style="Heading 2 for Appendix" if "Heading 2 for Appendix"
         in [s.name for s in doc.styles] else "Heading 2")
    lines = (D / "appendix_b_review_matrix.md").read_text(encoding="utf-8").split("\n")
    apx_caption = re.compile(r"^Table B[1-8]$")
    buffer: list[list[str]] = []
    prose: list[str] = []

    def flush_prose() -> None:
        if prose:
            para(doc, " ".join(prose), indent_first=Inches(0.5))
            prose.clear()

    for line in lines:
        t = line.strip()
        if t.startswith("|"):
            flush_prose()
            cells = [c.strip() for c in t.strip("|").split("|")]
            if not all(set(c) <= set("-: ") for c in cells):
                buffer.append(cells)
            continue
        if buffer:
            add_table(doc, buffer)
            buffer = []
        if not t:
            flush_prose()
            continue
        if t.startswith("# "):
            continue
        if apx_caption.match(t):
            flush_prose()
            para(doc, t, style="No Spacing", indent_first=Pt(0))
        elif t.startswith("*") and t.endswith("*") and t.count("*") == 2:
            flush_prose()
            para(doc, t, style="Table/Figure", indent_first=Pt(0))
        else:
            prose.append(t)
    flush_prose()
    if buffer:
        add_table(doc, buffer)


# ---------------------------------------------------------------------------------- main


def parse_front_matter() -> dict:
    text = (D / "front_matter.md").read_text(encoding="utf-8")
    parts = re.split(r"@([A-Z]+)@", text)[1:]
    fm = {parts[i].lower(): parts[i + 1].strip() for i in range(0, len(parts), 2)}
    return fm


def main(listings: dict | None = None) -> None:
    doc = Document(str(TEMPLATE))
    body = doc.element.body
    for child in list(body):
        if child.tag != qn("w:sectPr"):
            body.remove(child)

    fm = parse_front_matter()
    fm["listings"] = listings or {}
    front_matter(doc, fm)

    doc.add_section(WD_SECTION.NEW_PAGE)
    for heading, filename in CHAPTERS:
        add_markdown(doc, D / filename, heading)
    references(doc)
    appendices(doc)

    front, main_section = doc.sections[0], doc.sections[-1]
    page_numbering(front, "lowerRoman", 1, title_page=True)
    page_numbering(main_section, "decimal", 1, title_page=False)
    footer_page_number(front)
    footer_page_number(main_section)

    OUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUT)
    print(f"wrote {OUT}")
    print(f"  paragraphs {len(doc.paragraphs)}, tables {len(doc.tables)}, "
          f"figures {len(doc.inline_shapes)}, sections {len(doc.sections)}")


def single_chapter(index: int) -> Path:
    """Build one chapter on its own, in the template's styles, for supervisor review."""
    heading, filename = CHAPTERS[index - 1]
    doc = Document(str(TEMPLATE))
    body = doc.element.body
    for child in list(body):
        if child.tag != qn("w:sectPr"):
            body.remove(child)
    add_markdown(doc, D / filename, heading)
    if index == 2:
        # The review matrix answers a question raised on this chapter, so it travels with it.
        para(doc, "Appendices", style="Heading 1")
        appendix_b(doc)
    section = doc.sections[0]
    page_numbering(section, "decimal", 1, title_page=False)
    footer_page_number(section)
    out = D / f"Chapter_{['One','Two','Three','Four','Five'][index-1]}_for_review.docx"
    doc.save(out)
    print(f"wrote {out} ({len(doc.paragraphs)} paragraphs, {len(doc.tables)} tables, "
          f"{len(doc.inline_shapes)} figures)")
    return out


if __name__ == "__main__":
    if len(sys.argv) > 2 and sys.argv[1] == "--chapter":
        single_chapter(int(sys.argv[2]))
    else:
        main()
