"""Populate the three listings with the page numbers the document actually has.

There is no way to know a page number without laying the document out, so the document is
built, rendered, read back, and rebuilt with what the rendering reported. Adding listing
lines lengthens the front matter, which can move the pages of the listings that follow, so
this repeats until two consecutive passes agree.

Body page numbers cannot move, because the front matter is numbered in roman and the body
restarts at arabic one, so the loop converges in two or three passes.
"""
from __future__ import annotations
import re
import subprocess
import sys
from pathlib import Path

from docx import Document

NUMPAT = r"(?:\d+\.\d+|[A-Z]\.\d+)"
# The List of Tables covers the numbered chapter tables. Appendix tables carry their
# own B series and are reached through the appendix heading in the contents.
CHAPTER_NUM = r"\d+\.\d+"

sys.path.insert(0, str(Path(__file__).parent))
import build_dissertation as bd

DOCX = bd.OUT
WORK = Path("/tmp/paginate")
ROMAN = [(1000, "m"), (900, "cm"), (500, "d"), (400, "cd"), (100, "c"), (90, "xc"),
         (50, "l"), (40, "xl"), (10, "x"), (9, "ix"), (5, "v"), (4, "iv"), (1, "i")]


def roman(n: int) -> str:
    out = ""
    for value, sign in ROMAN:
        while n >= value:
            out += sign
            n -= value
    return out


def render() -> list[str]:
    WORK.mkdir(parents=True, exist_ok=True)
    for stale in WORK.glob("*"):
        stale.unlink()
    subprocess.run(["soffice", "-env:UserInstallation=file:///tmp/lo-paginate", "--headless",
                    "--norestore", "--convert-to", "pdf", "--outdir", str(WORK), str(DOCX)],
                   check=True, capture_output=True, timeout=900)
    pdf = next(WORK.glob("*.pdf"))
    txt = WORK / "doc.txt"
    subprocess.run(["pdftotext", "-layout", str(pdf), str(txt)], check=True, timeout=300)
    return txt.read_text(encoding="utf-8", errors="replace").split("\f")


def wanted() -> list[tuple[int, str, str]]:
    """Everything that needs a page number, in document order, as (level, text, kind)."""
    doc = Document(str(DOCX))
    items = []
    for p in doc.paragraphs:
        t = p.text.strip()
        if not t:
            continue
        style = p.style.name
        if style == "Section Title" and t != "ABSTRACT":
            items.append((0, t, "front"))
        elif style == "Section Title":
            items.append((0, t, "front"))
        elif style == "Heading 1":
            items.append((0, t, "toc"))
        elif style in ("Heading 2", "Heading 2 for Appendix"):
            # The template's appendix heading style is a Heading 2 variant, so without this
            # the appendices reach the document but never the table of contents.
            items.append((1, t, "toc"))
        elif style == "Heading 3":
            items.append((2, t, "toc"))
        elif re.fullmatch(rf"Table {CHAPTER_NUM}", t):
            items.append((0, t, "table"))
        elif style == "Table/Figure" and re.match(rf"Figure {NUMPAT}:", t):
            items.append((0, t, "figure"))
    return items


def locate(pages: list[str], items: list[tuple[int, str, str]]) -> dict[str, int]:
    """Find the first page carrying each item, scanning forward in document order.

    Front matter is located first and everything else only from the page after the last
    front-matter heading. Without that split, the second pass matches a chapter title
    against its own line in the table of contents, which puts the start of the body on
    the contents page and numbers the rest of the front matter in arabic.
    """
    found = {}

    def scan(subset, start):
        cursor = start
        for _, text, _ in subset:
            probe = re.sub(r"\s+", " ", text).strip().lower()[:58]
            for i in range(cursor, len(pages)):
                # Drop the listing lines themselves before matching. Each carries a dot
                # leader, and without this "LIST OF TABLES" matches its own entry on the
                # contents page rather than the heading on its own page.
                kept = [ln for ln in pages[i].split("\n") if "...." not in ln]
                body = re.sub(r"[ \t]+", " ", " ".join(kept)).lower()
                if probe in body:
                    found[text] = i + 1
                    cursor = i
                    break
            else:
                found[text] = cursor + 1
        return cursor

    front = [it for it in items if it[2] == "front"]
    rest = [it for it in items if it[2] != "front"]
    last_front = scan(front, 0)
    scan(rest, last_front + 1)
    return found


def shorten(caption: str, limit: int = 88) -> str:
    """One line per listing entry: keep the first sentence, and trim it if still too long."""
    # Captions now carry chapter-based numbers ("Table 3.10", "Figure 4.1:"), so the
    # number itself contains a full stop and cannot end the first sentence.
    head = re.match(rf"^((?:Table|Figure) {NUMPAT}[.:]\s*[^.]*)", caption)
    text = (head.group(1) if head else caption).strip().rstrip(".")
    if len(text) > limit:
        # Mark a trim rather than letting the listing quietly state a different title from
        # the one above the table. verify_document fails on any entry carrying this mark.
        text = text[:limit - 1].rsplit(" ", 1)[0] + "\u2026"
    return text


def build_listings(pages: list[str], items, found) -> dict:
    first_body = next((found[t] for _, t, k in items
                       if k == "toc" and t.lower().startswith("chapter one")), None)
    if first_body is None:
        raise SystemExit("could not find where the body starts")
    front_pages = first_body - 1

    def label(physical: int) -> str:
        return roman(physical) if physical <= front_pages else str(physical - front_pages)

    contents, tables, figures = [], [], []
    doc = Document(str(DOCX))
    paras = [p for p in doc.paragraphs]
    titles = {}
    for i, p in enumerate(paras):
        if re.fullmatch(rf"Table {CHAPTER_NUM}", p.text.strip()):
            nxt = paras[i + 1].text.strip() if i + 1 < len(paras) else ""
            titles[p.text.strip()] = nxt
    for level, text, kind in items:
        page = label(found[text])
        if kind in ("front", "toc"):
            contents.append((level, text, page))
        elif kind == "table":
            tables.append((0, shorten(f"{text}. {titles.get(text, '')}".rstrip(". ")), page))
        elif kind == "figure":
            figures.append((0, shorten(text), page))
    return {"contents": contents, "tables": tables, "figures": figures}


def main() -> None:
    previous = None
    for attempt in range(1, 5):
        pages = render()
        items = wanted()
        found = locate(pages, items)
        listings = build_listings(pages, items, found)
        signature = [(t, p) for rows in listings.values() for _, t, p in rows]
        print(f"pass {attempt}: {len(pages)} pages, {len(listings['contents'])} contents rows, "
              f"{len(listings['tables'])} tables, {len(listings['figures'])} figures")
        if signature == previous:
            print("page numbers are stable")
            return
        previous = signature
        bd.main(listings=listings)
    print("stopped after four passes without two agreeing", file=sys.stderr)
    sys.exit(1)


if __name__ == "__main__":
    main()
